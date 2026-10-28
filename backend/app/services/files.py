"""Spreadsheet upload analysis.

Supports ``.xlsx``, ``.xlsm``, ``.csv``, and ``.txt`` only — each validated and
exercised by tests. ``openpyxl`` is used for workbook *structure* extraction;
it is not a calculation engine, and nothing here executes macros or formulas.

Privacy: uploaded content is written to a private temp file, analyzed locally,
and deleted immediately. Only the structured analysis (sheet names, headers,
types, small samples) is retained, with an expiry. Nothing is sent to an AI
provider unless the user explicitly consents per request.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
import tempfile
import uuid
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import Settings, get_settings
from app.core.errors import AppError
from app.db.models import FileMeta
from app.schemas.api import (
    AnalysisRecommendation,
    ColumnProfile,
    SheetProfile,
    WorkbookAnalysis,
)

logger = logging.getLogger(__name__)

SUPPORTED_TYPES = {"xlsx": "xlsx", "xlsm": "xlsm", "csv": "csv", "txt": "txt"}
_ZIP_MAGIC = b"PK\x03\x04"

MAX_SCAN_ROWS = 5000  # bound on rows scanned per sheet
MAX_SAMPLE_ROWS = 5
MAX_DISTINCT_TRACKED = 1000
MAX_COLUMNS = 256


def _safe_name(filename: str) -> str:
    base = Path(filename).name
    cleaned = re.sub(r"[^A-Za-z0-9._ -]", "_", base)
    return cleaned[:200] or "upload"


def _validate_upload(filename: str, content: bytes, settings: Settings) -> str:
    """Validate extension, size, and magic bytes. Returns normalized type."""
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in SUPPORTED_TYPES:
        raise AppError(
            "unsupported_file_type",
            f"Unsupported file type '.{ext}'. Supported: .xlsx, .xlsm, .csv, .txt.",
            status_code=415,
        )
    if len(content) == 0:
        raise AppError("empty_file", "The uploaded file is empty.", status_code=400)
    if len(content) > settings.max_upload_bytes:
        raise AppError(
            "file_too_large",
            f"File exceeds the {settings.max_upload_mb} MB upload limit.",
            status_code=413,
        )
    if ext in ("xlsx", "xlsm"):
        if not content.startswith(_ZIP_MAGIC):
            raise AppError(
                "invalid_file_structure",
                "The file does not look like a valid Excel workbook (missing ZIP signature).",
                status_code=422,
            )
    else:
        # Text-based formats must be decodable as text.
        try:
            content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                content.decode("latin-1")
            except UnicodeDecodeError as exc:
                raise AppError(
                    "invalid_file_structure",
                    "The file could not be decoded as text.",
                    status_code=422,
                ) from exc
    return SUPPORTED_TYPES[ext]


def _infer_type(values: list[Any]) -> str:
    kinds: set[str] = set()
    for v in values:
        if v is None:
            continue
        if isinstance(v, bool):
            kinds.add("boolean")
        elif isinstance(v, (int, float)):
            kinds.add("number")
        elif isinstance(v, datetime):
            kinds.add("date")
        else:
            s = str(v).strip()
            if not s:
                continue
            if re.fullmatch(r"-?\d+(\.\d+)?", s):
                kinds.add("number")
            elif re.fullmatch(r"\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?", s):
                kinds.add("date")
            else:
                kinds.add("text")
    if not kinds:
        return "empty"
    if len(kinds) == 1:
        return kinds.pop()
    return "mixed"


def _profile_columns(headers: list[str], rows: list[list[Any]]) -> list[ColumnProfile]:
    profiles: list[ColumnProfile] = []
    for idx, header in enumerate(headers):
        column_values = [row[idx] if idx < len(row) else None for row in rows]
        non_empty = [v for v in column_values if v is not None and str(v).strip() != ""]
        distinct: set[str] = set()
        for v in non_empty:
            if len(distinct) <= MAX_DISTINCT_TRACKED:
                distinct.add(str(v))
        samples = []
        for v in non_empty[:5]:
            text = str(v)
            samples.append(text[:60])
        profiles.append(
            ColumnProfile(
                name=header or f"Column{idx + 1}",
                inferred_type=_infer_type(non_empty),  # type: ignore[arg-type]
                non_empty=len(non_empty),
                empty=len(rows) - len(non_empty),
                distinct_sampled=len(distinct),
                sample_values=samples,
            )
        )
    return profiles


def _profile_rows(
    sheet_name: str,
    headers: list[str],
    rows: list[list[Any]],
    total_rows_estimate: int | None,
    truncated: bool,
) -> SheetProfile:
    seen: set[tuple] = set()
    duplicates = 0
    for row in rows:
        key = tuple(str(v) if v is not None else "" for v in row)
        if key in seen:
            duplicates += 1
        else:
            seen.add(key)
    return SheetProfile(
        name=sheet_name,
        total_rows_estimate=total_rows_estimate,
        total_columns=len(headers),
        scanned_rows=len(rows),
        truncated=truncated,
        headers=headers,
        columns=_profile_columns(headers, rows),
        sample_rows=[
            [None if v is None else str(v)[:60] for v in row] for row in rows[:MAX_SAMPLE_ROWS]
        ],
        duplicate_rows_scanned=duplicates,
    )


def _analyze_workbook_xlsx(path: Path, file_type: str, notes: list[str]) -> list[SheetProfile]:
    # keep_vba=False: macros are never loaded or executed during analysis.
    try:
        wb = load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    except Exception as exc:
        raise AppError(
            "invalid_file_structure",
            "The workbook could not be opened. It may be corrupted, password "
            "protected, or not a real .xlsx/.xlsm file.",
            status_code=422,
        ) from exc
    if file_type == "xlsm":
        notes.append(
            "Macro-enabled workbook (.xlsm): analysis reads sheet structure and values "
            "only. VBA macros are NOT loaded, preserved, or executed."
        )

    profiles: list[SheetProfile] = []
    try:
        for ws in wb.worksheets[:50]:
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                headers: list[str] = []
                profiles.append(SheetProfile(name=ws.title, headers=[], total_columns=0))
                continue
            headers = [
                (str(h).strip() if h is not None and str(h).strip() else f"Column{i + 1}")
                for i, h in enumerate(header_row[:MAX_COLUMNS])
            ]
            scanned: list[list[Any]] = []
            truncated = False
            for row in rows_iter:
                if len(row) == 0 or all(v is None for v in row):
                    continue
                scanned.append(list(row[:MAX_COLUMNS]))
                if len(scanned) >= MAX_SCAN_ROWS:
                    truncated = True
                    break
            estimate = None
            try:
                if ws.max_row and ws.max_row > 0:
                    estimate = int(ws.max_row)
            except Exception:
                estimate = None
            profiles.append(_profile_rows(ws.title, headers, scanned, estimate, truncated))
    finally:
        wb.close()
    return profiles


def _analyze_text(
    path: Path, file_type: str, notes: list[str], label: str | None = None
) -> list[SheetProfile]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    else:  # pragma: no cover - latin-1 always decodes
        text = raw.decode("utf-8", errors="replace")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    all_rows: list[list[Any]] = []
    for i, row in enumerate(reader):
        all_rows.append(row)
        if i >= MAX_SCAN_ROWS:
            break
    truncated = len(all_rows) > MAX_SCAN_ROWS
    all_rows = all_rows[:MAX_SCAN_ROWS]
    if not all_rows:
        raise AppError("empty_file", "The file contains no rows.", status_code=400)
    headers = [(h.strip() or f"Column{i + 1}") for i, h in enumerate(all_rows[0][:MAX_COLUMNS])]
    body = [row[:MAX_COLUMNS] for row in all_rows[1:]]
    notes.append(
        f"Text data analyzed with delimiter '{delimiter}'. Header row taken from the first line."
    )
    return [_profile_rows(label or path.stem, headers, body, len(body) + 1, truncated)]


def _recommendations(sheets: list[SheetProfile]) -> list[AnalysisRecommendation]:
    recs: list[AnalysisRecommendation] = []
    for sheet in sheets:
        if sheet.duplicate_rows_scanned > 0:
            recs.append(
                AnalysisRecommendation(
                    area=f"{sheet.name}: duplicates",
                    recommendation=(
                        f"{sheet.duplicate_rows_scanned} duplicate row(s) detected among "
                        "the scanned rows. Ask sheetJEDI for a de-duplication solution "
                        "(Remove Duplicates / Power Query / script)."
                    ),
                )
            )
        empty_heavy = [
            c.name
            for c in sheet.columns
            if c.non_empty > 0 and c.empty > max(3, sheet.scanned_rows * 0.4)
        ]
        if empty_heavy:
            recs.append(
                AnalysisRecommendation(
                    area=f"{sheet.name}: blanks",
                    recommendation=(
                        f"Column(s) {', '.join(empty_heavy[:5])} have many blank cells. "
                        "Consider a cleaning plan (fill with default, or exclude from "
                        "aggregations)."
                    ),
                )
            )
        mixed = [c.name for c in sheet.columns if c.inferred_type == "mixed"]
        if mixed:
            recs.append(
                AnalysisRecommendation(
                    area=f"{sheet.name}: mixed types",
                    recommendation=(
                        f"Column(s) {', '.join(mixed[:5])} contain mixed data types "
                        "(e.g. numbers stored as text). Type conversion may be needed "
                        "before aggregation."
                    ),
                )
            )
        if sheet.truncated:
            recs.append(
                AnalysisRecommendation(
                    area=f"{sheet.name}: large sheet",
                    recommendation=(
                        f"Only the first {sheet.scanned_rows} rows were scanned for "
                        "performance. Statistics for this sheet are partial."
                    ),
                )
            )
    if not recs:
        recs.append(
            AnalysisRecommendation(
                area="general",
                recommendation=(
                    "No obvious data-quality issues detected in the scanned rows. "
                    "Describe the task you want to accomplish with this data."
                ),
            )
        )
    return recs


def purge_expired_files(db: Session) -> int:
    now = datetime.now(UTC)
    expired = db.scalars(
        select(FileMeta).where(FileMeta.expires_at.isnot(None), FileMeta.expires_at <= now)
    ).all()
    for meta in expired:
        db.delete(meta)
    if expired:
        db.commit()
    return len(expired)


async def analyze_upload(
    db: Session,
    filename: str,
    content: bytes,
    settings: Settings | None = None,
) -> WorkbookAnalysis:
    settings = settings or get_settings()
    purge_expired_files(db)
    file_type = _validate_upload(filename, content, settings)
    safe = _safe_name(filename)
    notes: list[str] = []

    # Content never persists: private temp file, deleted in `finally`.
    tmp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{file_type}", dir=None) as tmp:
            tmp.write(content)
            tmp_path = Path(tmp.name)
        if file_type in ("xlsx", "xlsm"):
            sheets = _analyze_workbook_xlsx(tmp_path, file_type, notes)
        else:
            # Label with the original (sanitized) filename, never the temp path.
            sheets = _analyze_text(tmp_path, file_type, notes, label=Path(safe).stem)
    finally:
        if tmp_path is not None:
            try:
                tmp_path.unlink(missing_ok=True)
            except OSError:  # pragma: no cover
                logger.warning("Could not delete temp upload file")

    notes.append(
        "Analysis performed locally with openpyxl/csv. Formula results were not "
        "calculated (cached values only) and no macros were executed."
    )

    analysis = WorkbookAnalysis(
        file_id=uuid.uuid4().hex,
        filename=safe,
        file_type=file_type,
        size_bytes=len(content),
        sheet_count=len(sheets),
        sheets=sheets,
        recommendations=_recommendations(sheets),
        notes=notes,
        analyzed_at=datetime.now(UTC),
    )

    meta = FileMeta(
        id=analysis.file_id,
        original_name=filename[:255],
        safe_name=safe,
        file_type=file_type,
        size_bytes=len(content),
        status="analyzed",
        analysis_json=analysis.model_dump_json(),
        expires_at=datetime.now(UTC) + timedelta(minutes=max(1, settings.upload_retention_minutes)),
    )
    db.add(meta)
    db.commit()
    return analysis


def delete_file(db: Session, file_id: str) -> None:
    meta = db.get(FileMeta, file_id)
    if meta is None:
        raise AppError(
            "file_not_found", "File analysis not found or already deleted.", status_code=404
        )
    db.delete(meta)
    db.commit()


def get_analysis_text(analysis_json: str, max_chars: int = 4000) -> str:
    """Compact text rendering of a stored analysis (for prompt context)."""
    try:
        data: dict[str, Any] = json.loads(analysis_json)
    except json.JSONDecodeError:
        return ""
    lines: list[str] = []
    for sheet in data.get("sheets", [])[:10]:
        headers = ", ".join(sheet.get("headers", [])[:30])
        lines.append(f"Sheet '{sheet.get('name')}': [{headers}]")
        for row in sheet.get("sample_rows", [])[:3]:
            lines.append("  " + ", ".join(str(v)[:40] for v in row[:10]))
    text = "\n".join(lines)
    return text[:max_chars]
